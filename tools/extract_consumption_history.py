"""Read-only extraction of the approved TD thuc layout. No recalculation or source edits."""
import openpyxl,json,sys,math
from pathlib import Path
from openpyxl.utils import column_index_from_string,get_column_letter
source=Path(sys.argv[1]);out=Path(sys.argv[2]);out.mkdir(parents=True,exist_ok=True)
w=openpyxl.load_workbook(source,data_only=False);cached=openpyxl.load_workbook(source,data_only=True)
s=w['TD thuc'];v=cached['TD thuc']
blocks=[(2014,'AJ'),(2015,'AX'),(2016,'BL'),(2017,'BZ'),(2018,'CO'),(2019,'EJ'),(2020,'EW'),(2021,'FO'),(2022,'GH'),(2023,'HA'),(2024,'HT'),(2025,'IM'),(2026,'JF')]
periods=[(2012,0,'D','ACTUAL'),(2013,0,'E','ACTUAL')]
for year,start in blocks:
 for month in range(1,13):periods.append((year,month,get_column_letter(column_index_from_string(start)+month-1),'ACTUAL'))
for year,start in [(2018,'DC'),(2019,'DU')]:
 for month in range(1,13):periods.append((year,month,get_column_letter(column_index_from_string(start)+month-1),'PLAN'))
records=[];issues=[];empty=0
for base in range(21,131,11):
 code=str(s.cell(base,2).value).strip()
 if code=='VTCM':continue
 code='VTB' if code=='VBD' else code
 for year,month,col,scenario in periods:
  cells=[f'{col}{base+k}' for k in range(1,6)];vals=[v[a].value for a in cells]
  if all(x is None or x=='' for x in vals):empty+=1;continue
  record=dict(marketCode=code,year=year,month=month,scenario=scenario,inputMode='AMOUNTS',status='DRAFT',source=f'{source.name} | TD thuc!{cells[0]}:{cells[-1]}',note='Imported historical values; unit thousand USD. '+('Uses saved Excel formula results; external links not refreshed.' if any(s[a].data_type=='f' for a in cells) else ''))
  errors=[]
  for key,a,value in zip(['total','voice','sms','data','vas'],cells,vals):
   if value is None or value=='':record[key]=''
   elif isinstance(value,(int,float)) and not isinstance(value,bool) and math.isfinite(value) and value>=0:record[key]=value
   else:errors.append(f'{a}: invalid/missing cached numeric value {value!r}');record[key]=''
  if errors:issues.append(dict(record=record,errors=errors));continue
  total=record['total'];parts=[record[k] for k in ['voice','sms','data','vas']]
  if total!='' and all(x!='' for x in parts) and abs(sum(parts)-total)>max(.001,abs(total)*.0001):errors.append(f'Total {total} differs from sum {sum(parts)}')
  elif total!='' and sum(x for x in parts if x!='')>total+max(.001,abs(total)*.0001):errors.append('Known components exceed total')
  if errors:
   issues.append(dict(record=record,errors=errors,imported=True))
   record['note']+=' Original total retained; service breakdown needs reconciliation.'
  records.append(record)
(out/'PR06_history_import.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
(out/'PR06_history_review.json').write_text(json.dumps(issues,ensure_ascii=False,indent=2),encoding='utf-8')
from collections import Counter
summary=dict(ready=len(records),review=len(issues),importedFlagged=sum(bool(x.get('imported')) for x in issues),excluded=sum(not x.get('imported',False) for x in issues),empty=empty,byYear=dict(Counter(str(r['year']) for r in records)),byMarket=dict(Counter(r['marketCode'] for r in records)),policies=['thousand USD','VBD to VTB','VTCM excluded','2026 actual','cached values only','original total wins; mismatches imported as Draft and flagged','annual-only 2012/2013 kept as month 0'])
(out/'PR06_history_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(summary,ensure_ascii=False));print('Review samples:',json.dumps(issues[:3],ensure_ascii=False))
