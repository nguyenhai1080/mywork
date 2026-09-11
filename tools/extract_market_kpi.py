import openpyxl,json,sys
from pathlib import Path
w=openpyxl.load_workbook(sys.argv[1],data_only=True,read_only=True);s=w['TH thi truong']
fields={7:'sites',8:'registered',9:'active',10:'population',18:'marketShare',19:'arpuActual',20:'arpuFixed',21:'arpu4g',22:'exchangeRate',24:'dataSubsShare',25:'data4gShare',28:'dou',29:'dou4g',30:'dataPrice',31:'dataCost',35:'mouOutgoing',36:'voicePrice',37:'voiceCost'}
rows=[]
for col in list(range(3,12))+[13]:
 code=s.cell(1,col).value;code='VTB' if code=='VBD' else code
 values={k:s.cell(row,col).value for row,k in fields.items()}
 for k in ['marketShare','dataSubsShare','data4gShare']:
  if isinstance(values[k],(int,float)):values[k]*=100
 for k,v in values.items():
  if v is not None and not isinstance(v,(int,float)):raise ValueError((code,k,v))
 rows.append(dict(marketCode=code,year=s['B1'].value,month=s['A1'].value,values=values,references={'sites':'2025','population':'2025','marketShare':'03/2026','dataCost':'2025','voiceCost':'2025'},source='TD thuc.xlsx | TH thi truong | column '+openpyxl.utils.get_column_letter(col),note='Cached workbook values; external links not refreshed. DOU, MOU and zero values retained as supplied.',revision=0))
for r in rows:
 if r['marketCode']=='VTT':r['values']['originalDataShare']=s['M26'].value*100
# Preserve original composition snapshots from the report, separate from consumption amounts.
services=['Voice','SMS','Data','VAS']
for first,year,month in [(51,int(s['B1'].value),int(s['A1'].value)),(57,2025,12),(63,2024,12),(69,2023,12),(75,2022,12),(81,2021,12),(87,2020,12),(93,2019,12),(100,2018,12)]:
 for col in range(3,14):
  code=s.cell(first-1,col).value
  if not code:continue
  code=str(code).strip();code='VTB' if code=='VBD' else code
  shares={ 'original'+key+'Share':s.cell(first+i,col).value*100 if isinstance(s.cell(first+i,col).value,(int,float)) else None for i,key in enumerate(services)}
  if not any(v is not None for v in shares.values()):continue
  existing=next((r for r in rows if r['marketCode']==code and r['year']==year and r['month']==month),None)
  if existing:existing['values'].update(shares)
  else:rows.append(dict(marketCode=code,year=year,month=month,values=shares,references={},source='TD thuc.xlsx | TH thi truong!'+openpyxl.utils.get_column_letter(col)+str(first)+':'+openpyxl.utils.get_column_letter(col)+str(first+3),note='Original report composition snapshot; not a replacement for consumption amounts.',revision=0))
Path(sys.argv[2]).write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8');print(len(rows))
